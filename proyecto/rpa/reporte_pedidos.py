"""RPA: entra al frontend (React) como un usuario, lee la lista de pedidos del día y genera reporte_diario.xlsx.

Variables de entorno:
  RPA_USUARIO, RPA_CLAVE   usuario de prueba del panel (obligatorias; nunca van escritas en el código)
  FRONT_URL                dirección del frontend (por defecto http://localhost:5173)
  HEADLESS                 "1" = sin ventana (por defecto "0": se ve el navegador)
  TODOS                    "1" = incluir todos los pedidos, no solo los de hoy

Uso:  python reporte_pedidos.py
"""
import logging
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import expect, sync_playwright

AQUI = Path(__file__).parent
CAPTURAS = AQUI / "capturas"
SALIDA = AQUI / "reporte_diario.xlsx"

FRONT_URL = os.environ.get("FRONT_URL", "http://localhost:5173").rstrip("/")
USUARIO = os.environ.get("RPA_USUARIO", "")
CLAVE = os.environ.get("RPA_CLAVE", "")
HEADLESS = os.environ.get("HEADLESS", "0") == "1"
TODOS = os.environ.get("TODOS", "0") == "1"

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")
log = logging.getLogger("rpa")


def entrar(pagina):
    """Inicia sesión igual que una persona: escribe usuario y clave y pulsa Entrar."""
    log.info("Abriendo %s/login", FRONT_URL)
    pagina.goto(f"{FRONT_URL}/login")
    pagina.get_by_label("Usuario").fill(USUARIO)
    pagina.get_by_label("Clave").fill(CLAVE)
    pagina.get_by_role("button", name="Entrar").click()
    try:
        pagina.wait_for_url("**/pedidos", timeout=10_000)
    except PlaywrightTimeout:
        # Si el login falla, la URL no cambia. Se distingue "clave mala" de "la página no responde"
        if pagina.get_by_text("Usuario o clave incorrectos").is_visible():
            raise ValueError("Login rechazado: usuario o clave incorrectos") from None
        raise
    log.info("Sesión iniciada")


def leer_pagina_actual(pagina):
    """Devuelve las filas de la tabla visible: (numero, cliente, fecha_texto, estado, total)."""
    # Se espera a que termine de cargar: el contenedor deja de estar "ocupado" (aria-busy) o aparece el mensaje de vacío.
    # Sin esto se leería la tabla vieja o el texto "Cargando pedidos…" (que también usa la clase .vacio).
    pagina.wait_for_selector(
        '.tabla-contenedor[aria-busy="false"] tbody tr, .vacio:not(:has-text("Cargando"))', timeout=10_000
    )
    filas = []
    for tr in pagina.locator("table tbody tr").all():
        celdas = [c.inner_text().strip() for c in tr.locator("td").all()]
        numero, cliente, fecha, estado, total = celdas[:5]
        filas.append({
            "pedido": int(numero.lstrip("#")),
            "cliente": cliente,
            "fecha_texto": fecha,
            "estado": estado,
            # "$ 45.000" -> 45000 (el formato es COP sin decimales)
            "total": int(re.sub(r"\D", "", total) or 0),
        })
    return filas


def leer_todos_los_pedidos(pagina):
    """Pone 50 por página y recorre todas las páginas con el botón Siguiente."""
    pagina.locator("#filtro-limite").select_option("50")
    resultado = []
    while True:
        resultado += leer_pagina_actual(pagina)
        siguiente = pagina.get_by_role("button", name="Siguiente")
        if siguiente.count() == 0 or siguiente.is_disabled():
            return resultado
        indicador = pagina.locator("nav[aria-label='Paginación'] p")
        antes = indicador.inner_text()
        siguiente.click()
        # Espera a que el indicador "Página X de Y" cambie: así sabemos que cargó la página nueva
        expect(indicador).not_to_have_text(antes, timeout=10_000)


def generar_excel(filas, titulo_dia):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pedidos"
    encabezados = ["Pedido", "Cliente", "Fecha", "Estado", "Total (COP)"]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2F4F4F")
        celda.alignment = Alignment(horizontal="center")

    for f in filas:
        hoja.append([f["pedido"], f["cliente"], f["fecha_texto"], f["estado"], f["total"]])

    total = sum(f["total"] for f in filas)
    hoja.append(["", "", "", "TOTAL", total])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)

    for fila in hoja.iter_rows(min_row=2, min_col=5, max_col=5):
        fila[0].number_format = '"$"#,##0'
    for letra, ancho in zip("ABCDE", (10, 28, 26, 14, 16)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"

    # Segunda hoja: resumen por estado
    resumen = libro.create_sheet("Resumen por estado")
    resumen.append(["Estado", "Cantidad", "Total (COP)"])
    for celda in resumen[1]:
        celda.font = Font(bold=True)
    por_estado = defaultdict(lambda: [0, 0])
    for f in filas:
        por_estado[f["estado"]][0] += 1
        por_estado[f["estado"]][1] += f["total"]
    for estado, (cantidad, suma) in sorted(por_estado.items()):
        resumen.append([estado, cantidad, suma])
    for fila in resumen.iter_rows(min_row=2, min_col=3, max_col=3):
        fila[0].number_format = '"$"#,##0'
    resumen["E1"] = f"Reporte del {titulo_dia}"
    for letra, ancho in zip("ABC", (14, 10, 16)):
        resumen.column_dimensions[letra].width = ancho

    libro.save(SALIDA)


def main() -> int:
    if not USUARIO or not CLAVE:
        print("Faltan RPA_USUARIO y RPA_CLAVE. En PowerShell:\n"
              '  $env:RPA_USUARIO="admin"; $env:RPA_CLAVE="tu_clave"', file=sys.stderr)
        return 2

    CAPTURAS.mkdir(exist_ok=True)
    with sync_playwright() as p:
        navegador = p.chromium.launch(headless=HEADLESS)
        # Misma zona horaria e idioma que el frontend muestra: así "hoy" significa lo mismo para el bot y la persona
        contexto = navegador.new_context(locale="es-CO", timezone_id="America/Bogota", viewport={"width": 1280, "height": 900})
        pagina = contexto.new_page()
        try:
            entrar(pagina)
            todos = leer_todos_los_pedidos(pagina)
            pagina.screenshot(path=str(CAPTURAS / "pedidos_leidos.png"), full_page=True)

            # "Hoy" se calcula con el formato del propio navegador y se compara con la parte de fecha (antes de la coma) de cada fila
            hoy = pagina.evaluate("new Date().toLocaleDateString('es-CO', { dateStyle: 'medium' })")
            del_dia = todos if TODOS else [f for f in todos if f["fecha_texto"].split(",")[0].strip() == hoy]
            log.info("Leídos %d pedidos; del día (%s): %d", len(todos), hoy, len(del_dia))

            generar_excel(del_dia, hoy)
            print(f"OK: {SALIDA.name} con {len(del_dia)} pedidos (total leído en pantalla: {len(todos)}).")
            return 0
        except (PlaywrightError, AssertionError, ValueError, IndexError) as error:
            destino = CAPTURAS / "error_reporte.png"
            pagina.screenshot(path=str(destino), full_page=True)
            log.error("Falló el bot: %s", error)
            print(f"ERROR: {error}\nCaptura guardada en {destino}", file=sys.stderr)
            return 1
        finally:
            navegador.close()


if __name__ == "__main__":
    sys.exit(main())

import argparse
import csv
import logging
from datetime import datetime

from openpyxl import Workbook

logger = logging.getLogger(__name__)


def leer_csv(ruta):
    with open(ruta, "r", encoding="utf-8", newline="") as archivo:
        return list(csv.DictReader(archivo))


def limpiar_datos(filas):
    """Quita duplicados, normaliza la ciudad y descarta totales o fechas inválidos."""
    total_inicial = len(filas)

    vistos = set()
    sin_duplicados = []
    for fila in filas:
        clave = tuple(fila.values())
        if clave not in vistos:
            vistos.add(clave)
            sin_duplicados.append(fila)
    logger.info("Duplicados eliminados: %d", total_inicial - len(sin_duplicados))

    limpias = []
    for fila in sin_duplicados:
        try:
            total = float(fila["total"])
            fecha = datetime.strptime(fila["fecha"], "%Y-%m-%d")
        except (ValueError, TypeError):
            continue                     # total vacío o fecha mal escrita
        limpias.append({
            "id": fila["id"],
            "fecha": fecha,
            "ciudad": fila["ciudad"].strip().title(),
            "producto": fila["producto"].strip(),
            "total": total,
        })

    logger.info(
        "Filas descartadas por total o fecha inválidos: %d",
        len(sin_duplicados) - len(limpias),
    )
    logger.info("Filas limpias: %d de %d", len(limpias), total_inicial)
    return limpias


def calcular_analisis(filas):
    """Devuelve total por mes, por ciudad y el top 5 de productos."""
    por_mes, por_ciudad, por_producto = {}, {}, {}
    for fila in filas:
        mes = fila["fecha"].strftime("%Y-%m")
        por_mes[mes] = por_mes.get(mes, 0) + fila["total"]
        por_ciudad[fila["ciudad"]] = por_ciudad.get(fila["ciudad"], 0) + fila["total"]
        por_producto[fila["producto"]] = por_producto.get(fila["producto"], 0) + fila["total"]

    meses = sorted(por_mes.items())
    ciudades = sorted(por_ciudad.items())
    top5 = sorted(por_producto.items(), key=lambda par: par[1], reverse=True)[:5]
    return meses, ciudades, top5


def generar_excel(meses, ciudades, top5, ruta_salida):
    libro = Workbook()
    libro.remove(libro.active)           # quita la hoja vacía que crea Workbook()

    hojas = [
        ("Por mes", ["mes", "total"], meses),
        ("Por ciudad", ["ciudad", "total"], ciudades),
        ("Top 5 productos", ["producto", "total"], top5),
    ]
    for nombre, encabezado, filas in hojas:
        hoja = libro.create_sheet(nombre)
        hoja.append(encabezado)
        for fila in filas:
            hoja.append(list(fila))

    libro.save(ruta_salida)
    logger.info("Excel generado: %s", ruta_salida)


def main():
    parser = argparse.ArgumentParser(description="Genera un reporte de ventas en Excel")
    parser.add_argument("--entrada", required=True, help="CSV de ventas")
    parser.add_argument("--salida", required=True, help="Excel de salida (.xlsx)")
    parser.add_argument("--verbose", action="store_true", help="Muestra más detalle")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )

    try:
        filas = leer_csv(args.entrada)
    except FileNotFoundError:
        logger.error("No existe el archivo de entrada: %s", args.entrada)
        raise SystemExit(1)

    logger.info("Leídas %d filas de %s", len(filas), args.entrada)

    filas = limpiar_datos(filas)
    meses, ciudades, top5 = calcular_analisis(filas)
    generar_excel(meses, ciudades, top5, args.salida)


if __name__ == "__main__":
    main()
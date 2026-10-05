from openpyxl import load_workbook

libro = load_workbook("reporte.xlsx")

for hoja in libro.worksheets:
    print(f"\n--- {hoja.title} ({hoja.max_row - 1} filas de datos) ---")
    for fila in hoja.iter_rows(values_only=True):
        print(fila)